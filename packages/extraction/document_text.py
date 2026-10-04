"""Bank-guarantee document text loading for extraction."""

from __future__ import annotations

import io
import tempfile
from pathlib import Path
from typing import Any, cast


def load_document_text(path: Path) -> str:
    """Return plain text suitable for the BG extraction prompt.

    Supports .msg (Outlook) with .docx attachments, .docx, and plain text.
    PDF/image bytes are not converted here — callers should use a vision path.
    """

    suffix = path.suffix.lower()
    if suffix == ".msg":
        return _text_from_msg(path)
    return load_document_text_from_bytes(path.read_bytes(), suffix=suffix)


def load_document_text_from_bytes(data: bytes, *, suffix: str = ".txt") -> str:
    """Load extractable text from bytes. ``suffix`` includes the leading dot."""

    normalized = suffix.lower() if suffix.startswith(".") else f".{suffix.lower()}"
    if normalized == ".msg":
        with tempfile.NamedTemporaryFile(suffix=".msg") as handle:
            handle.write(data)
            handle.flush()
            return _text_from_msg(Path(handle.name))
    if normalized == ".docx":
        return _text_from_docx_bytes(data)
    if normalized in {".txt", ".md", ".html", ".htm", ""}:
        return data.decode("utf-8", errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError(
            f"unsupported document type for text extraction: {normalized or 'unknown'}"
        ) from exc


def pages_from_text(document_text: str) -> list[dict[str, object]]:
    """Split into page dicts. Flat loaders yield one chunk with page=null."""

    text = document_text.strip()
    if not text:
        return []
    # No page markers from current loaders — never invent a page number.
    return [{"page": None, "text": text}]


def _text_from_msg(path: Path) -> str:
    import extract_msg

    msg = extract_msg.Message(str(path))  # type: ignore[no-untyped-call]
    parts: list[str] = []
    subject = (msg.subject or "").strip()
    if subject:
        parts.append(f"Subject: {subject}")
    body = (msg.body or "").strip()
    if body:
        parts.append(body)
    for attachment in msg.attachments:
        name = attachment.longFilename or attachment.shortFilename or "attachment"
        data = cast(Any, attachment.data)
        if not data:
            continue
        payload = bytes(data)
        lower = str(name).lower()
        if lower.endswith(".docx"):
            parts.append(f"--- attachment: {name} ---")
            parts.append(_text_from_docx_bytes(payload))
        elif lower.endswith((".txt", ".md")):
            parts.append(f"--- attachment: {name} ---")
            parts.append(payload.decode("utf-8", errors="replace"))
    text = "\n\n".join(parts).strip()
    if not text:
        raise ValueError(f"no extractable text in {path.name}")
    return text


def _text_from_docx_bytes(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [
                c.text.strip().replace("\n", " ") for c in row.cells if c.text.strip()
            ]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)
